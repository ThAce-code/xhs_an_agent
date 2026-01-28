from __future__ import annotations

import csv
import datetime as _dt
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any


_URL_RE = re.compile(r"https?://\S+")


def slugify(text: str, *, max_len: int = 40) -> str:
    s = re.sub(r"\s+", "_", text.strip())
    s = re.sub(r"[^0-9A-Za-z_\u4e00-\u9fff-]+", "", s)
    return (s[:max_len] or "query").strip("_")


def now_stamp() -> str:
    return _dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def extract_text(content: Any) -> str:
    """Best-effort extraction of human-readable text from model/tool outputs."""

    if content is None:
        return ""
    if isinstance(content, str):
        return content

    # Gemini/LangChain sometimes returns a list of parts: [{'type':'text','text':'...','extras':...}, ...]
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                txt = item.get("text") or item.get("content")
                if isinstance(txt, str):
                    parts.append(txt)
        if parts:
            return "\n".join(parts)
        return str(content)

    if isinstance(content, dict):
        txt = content.get("text") or content.get("content")
        if isinstance(txt, str):
            return txt
        return str(content)

    maybe = getattr(content, "content", None)
    if maybe is not None and maybe is not content:
        return extract_text(maybe)

    return str(content)


@dataclass
class ParsedReport:
    why_hot: dict[str, Any]
    structure: dict[str, Any]
    ideas: dict[str, Any]
    sources: list[dict[str, Any]]
    raw: str


@dataclass
class RewriteReport:
    drafts: list[dict[str, Any]]
    title_bank: list[dict[str, Any]]
    hooks: list[dict[str, Any]]
    hashtags: list[dict[str, Any]]
    cta: list[dict[str, Any]]
    compliance_notes: list[dict[str, Any]]
    raw: str


@dataclass
class CoverReport:
    cover_concepts: list[dict[str, Any]]
    shotlist: list[dict[str, Any]]
    canva_recipe: list[dict[str, Any]]
    image_prompt: dict[str, Any] | None
    raw: str


@dataclass
class RadarReport:
    niches: list[dict[str, Any]]
    top3: list[dict[str, Any]]
    how_to_validate: list[dict[str, Any]]
    sources: list[dict[str, Any]]
    raw: str


def _strip_code_fence(text: str) -> str:
    s = text.strip()
    if s.startswith("```"):
        s = re.sub(r"^```[a-zA-Z0-9_-]*\s*", "", s)
        s = re.sub(r"```\s*$", "", s)
    return s.strip()


def _slice_first_json_object(text: str) -> str | None:
    s = _strip_code_fence(text)
    if "{" not in s:
        return None
    start = s.find("{")
    end = s.rfind("}")
    if start == -1 or end == -1 or end <= start:
        return None
    return s[start : end + 1].strip()


def _ensure_list(v: Any) -> list[str]:
    if v is None:
        return []
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    if isinstance(v, str):
        return [v.strip()] if v.strip() else []
    return [str(v).strip()] if str(v).strip() else []


def _ensure_cites(v: Any) -> list[int]:
    if v is None:
        return []
    if isinstance(v, list):
        out: list[int] = []
        for x in v:
            try:
                out.append(int(x))
            except Exception:
                continue
        return out[:3]
    if isinstance(v, str):
        return [int(n) for n in re.findall(r"\d+", v)][:3]
    try:
        return [int(v)][:3]
    except Exception:
        return []


def _ensure_items(v: Any) -> list[dict[str, Any]]:
    """Normalize list-like fields into [{text, cites}] for stable exports."""

    def _norm_item(x: Any) -> dict[str, Any] | None:
        if x is None:
            return None
        if isinstance(x, dict):
            text = x.get("text") or x.get("content") or ""
            cites = x.get("cites") or x.get("citations") or []
            if isinstance(cites, list):
                out_cites: list[int] = []
                for c in cites:
                    try:
                        out_cites.append(int(c))
                    except Exception:
                        continue
            elif isinstance(cites, str):
                out_cites = [int(n) for n in re.findall(r"\d+", cites)]
            else:
                out_cites = []
            t = str(text).strip()
            if not t:
                return None
            return {"text": t, "cites": out_cites[:3]}
        t = str(x).strip()
        if not t:
            return None
        return {"text": t, "cites": []}

    if v is None:
        return []
    if isinstance(v, list):
        out: list[dict[str, Any]] = []
        for x in v:
            it = _norm_item(x)
            if it is not None:
                out.append(it)
        return out
    it = _norm_item(v)
    return [it] if it is not None else []


def _ensure_item(v: Any) -> dict[str, Any] | None:
    items = _ensure_items(v)
    return items[0] if items else None


def _ensure_drafts(v: Any) -> list[dict[str, Any]]:
    if not isinstance(v, list):
        return []
    out: list[dict[str, Any]] = []
    for x in v:
        if not isinstance(x, dict):
            continue
        title = x.get("title")
        body = x.get("body")
        variant = x.get("variant") or x.get("name") or ""
        if not isinstance(title, str) or not title.strip():
            continue
        if not isinstance(body, str) or not body.strip():
            continue
        out.append(
            {
                "variant": str(variant).strip() or "",
                "title": title.strip(),
                "body": body.strip(),
                "cites": _ensure_cites(x.get("cites")),
            }
        )
    return out


def _normalize_sources(v: Any, *, raw_text: str) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []

    if isinstance(v, list):
        for x in v:
            if isinstance(x, dict):
                url = x.get("url") or x.get("link")
                if not isinstance(url, str) or not url.strip():
                    continue
                sid = x.get("id")
                try:
                    sid_int = int(sid) if sid is not None else len(out) + 1
                except Exception:
                    sid_int = len(out) + 1
                out.append(
                    {
                        "id": sid_int,
                        "url": url.strip(),
                        "title": (x.get("title") if isinstance(x.get("title"), str) else "") or "",
                    }
                )
            elif isinstance(x, str) and x.strip():
                out.append({"id": len(out) + 1, "url": x.strip(), "title": ""})

    # Best-effort: scrape URLs from raw text and append if missing.
    for u in _URL_RE.findall(raw_text):
        u = u.rstrip(")],.，。;")
        if not any(s.get("url") == u for s in out):
            # Preserve existing id space if present.
            max_id = 0
            for s in out:
                try:
                    max_id = max(max_id, int(s.get("id") or 0))
                except Exception:
                    continue
            out.append({"id": max_id + 1, "url": u, "title": ""})

    # De-dupe by URL, keep first (lowest id).
    seen: set[str] = set()
    dedup: list[dict[str, Any]] = []
    for s in out:
        url = s.get("url")
        if not isinstance(url, str) or not url.strip() or url in seen:
            continue
        seen.add(url)
        dedup.append(s)

    # Stable ordering for exports (keep original ids if provided).
    def _id_key(x: dict[str, Any]) -> tuple[int, str]:
        try:
            return (int(x.get("id") or 0), str(x.get("url") or ""))
        except Exception:
            return (0, str(x.get("url") or ""))

    dedup.sort(key=_id_key)
    return dedup


def parse_report(text: str) -> ParsedReport:
    raw = text.strip()

    # Preferred: JSON schema produced by SYSTEM_PROMPT
    json_str = _slice_first_json_object(raw)
    if json_str is not None:
        try:
            data = json.loads(json_str)
            if isinstance(data, dict) and ("why_hot" in data or "ideas" in data or "structure" in data):
                why_hot = data.get("why_hot") if isinstance(data.get("why_hot"), dict) else {}
                structure = data.get("structure") if isinstance(data.get("structure"), dict) else {}
                ideas = data.get("ideas") if isinstance(data.get("ideas"), dict) else {}

                sources = _normalize_sources(data.get("sources"), raw_text=raw)

                return ParsedReport(
                    why_hot=why_hot,
                    structure=structure,
                    ideas=ideas,
                    sources=sources,
                    raw=raw,
                )
        except Exception:
            pass

    # Fallback: legacy sectioned text. We map it into the new structure.
    pattern = re.compile(r"【(?P<h>热门赛道|爆款关键词|可模仿选题|数据来源)】")
    matches = list(pattern.finditer(raw))
    sections: dict[str, str] = {}
    if matches:
        for i, m in enumerate(matches):
            h = m.group("h")
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(raw)
            sections[h] = raw[start:end].strip()

    def _bullets(block: str) -> list[str]:
        items: list[str] = []
        for line in block.splitlines():
            s = line.strip()
            if not s:
                continue
            s = re.sub(r"^[\*\-\u2022]\s+", "", s)
            items.append(s)
        return items

    tracks = _bullets(sections.get("热门赛道", ""))
    keywords = _bullets(sections.get("爆款关键词", ""))

    why_hot = {
        "core_pain_points": tracks,
        "emotional_hooks": [],
        "persona": [],
    }
    structure = {
        "title_templates": [],
        "visual_style": "",
        "seo_keywords": keywords,
    }
    ideas = {
        "follow": _bullets(sections.get("可模仿选题", "")),
        "reverse": [],
        "upgrade": [],
    }

    src_block = sections.get("数据来源", "") or raw
    sources = _normalize_sources([u.rstrip(")],.，。;") for u in _URL_RE.findall(src_block)], raw_text=src_block)

    return ParsedReport(
        why_hot=why_hot,
        structure=structure,
        ideas=ideas,
        sources=sources,
        raw=raw,
    )


def _parse_json_dict(text: str) -> dict[str, Any] | None:
    json_str = _slice_first_json_object(text.strip())
    if json_str is None:
        return None
    try:
        data = json.loads(json_str)
    except Exception:
        return None
    return data if isinstance(data, dict) else None


def parse_rewrite(text: str) -> RewriteReport:
    raw = text.strip()
    data = _parse_json_dict(raw) or {}

    drafts = _ensure_drafts(data.get("drafts"))
    title_bank = _ensure_items(data.get("title_bank"))
    hooks = _ensure_items(data.get("hooks"))
    hashtags = _ensure_items(data.get("hashtags"))
    cta = _ensure_items(data.get("cta"))
    compliance_notes = _ensure_items(data.get("compliance_notes"))

    return RewriteReport(
        drafts=drafts,
        title_bank=title_bank,
        hooks=hooks,
        hashtags=hashtags,
        cta=cta,
        compliance_notes=compliance_notes,
        raw=raw,
    )


def parse_cover(text: str) -> CoverReport:
    raw = text.strip()
    data = _parse_json_dict(raw) or {}

    cover_concepts: list[dict[str, Any]] = []
    for x in data.get("cover_concepts") if isinstance(data.get("cover_concepts"), list) else []:
        if not isinstance(x, dict):
            continue
        headline = x.get("headline")
        subheadline = x.get("subheadline")
        if not isinstance(headline, str) or not headline.strip():
            continue
        if not isinstance(subheadline, str) or not subheadline.strip():
            continue
        cover_concepts.append(
            {
                "name": str(x.get("name") or "").strip(),
                "headline": headline.strip(),
                "subheadline": subheadline.strip(),
                "mood": _ensure_list(x.get("mood")),
                "composition": str(x.get("composition") or "").strip(),
                "lighting": str(x.get("lighting") or "").strip(),
                "props": _ensure_list(x.get("props")),
                "color_palette": _ensure_list(x.get("color_palette")),
                "typography": str(x.get("typography") or "").strip(),
                "post_process": str(x.get("post_process") or "").strip(),
                "cites": _ensure_cites(x.get("cites")),
            }
        )

    shotlist = _ensure_items(data.get("shotlist"))
    canva_recipe = _ensure_items(data.get("canva_recipe"))
    image_prompt = _ensure_item(data.get("image_prompt"))

    return CoverReport(
        cover_concepts=cover_concepts,
        shotlist=shotlist,
        canva_recipe=canva_recipe,
        image_prompt=image_prompt,
        raw=raw,
    )


def parse_radar(text: str) -> RadarReport | None:
    """Parse the '账号方向雷达' JSON schema."""

    raw = text.strip()
    data = _parse_json_dict(raw)
    if not isinstance(data, dict):
        return None
    if not isinstance(data.get("niches"), list):
        return None

    niches: list[dict[str, Any]] = []
    for x in data.get("niches") or []:
        if not isinstance(x, dict):
            continue
        name = x.get("name")
        if not isinstance(name, str) or not name.strip():
            continue
        niches.append(x)

    top3_in = data.get("top3")
    top3: list[dict[str, Any]] = []
    if isinstance(top3_in, list):
        for x in top3_in:
            if isinstance(x, dict) and isinstance(x.get("name"), str) and str(x.get("name") or "").strip():
                top3.append(x)

    how_to_validate = _ensure_items(data.get("how_to_validate"))
    sources = _normalize_sources(data.get("sources"), raw_text=raw)

    return RadarReport(
        niches=niches,
        top3=top3,
        how_to_validate=how_to_validate,
        sources=sources,
        raw=raw,
    )


def format_report_text(report: ParsedReport) -> str:
    """Human-friendly output for terminal (Markdown-like)."""

    why = report.why_hot or {}
    st = report.structure or {}
    ideas = report.ideas or {}

    lines: list[str] = []

    lines.append("## Step 1: 流量归因分析 (Why is it hot?)")
    lines.append("- 核心痛点")
    for it in _ensure_items(why.get("core_pain_points")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("- 情绪钩子")
    for it in _ensure_items(why.get("emotional_hooks")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("- 人群画像")
    for it in _ensure_items(why.get("persona")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("")

    lines.append("## Step 2: 爆款元素拆解 (Structure)")
    lines.append("- 标题公式")
    for it in _ensure_items(st.get("title_templates")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("- 视觉风格")
    vs_it = _ensure_item(st.get("visual_style"))
    if vs_it is not None:
        cites = vs_it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {vs_it.get('text','')}{suffix}")
    lines.append("- 关键词布局")
    for it in _ensure_items(st.get("seo_keywords")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("")

    lines.append("## Step 3: 差异化选题建议 (How to copy & win?)")
    lines.append("- 跟风型")
    for it in _ensure_items(ideas.get("follow")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("- 反向型")
    for it in _ensure_items(ideas.get("reverse")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("- 升级型")
    for it in _ensure_items(ideas.get("upgrade")):
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"  - {it.get('text','')}{suffix}")
    lines.append("")

    lines.append("## 数据来源")
    for s in report.sources:
        sid = s.get("id")
        url = s.get("url")
        title = s.get("title") or ""
        if not isinstance(url, str) or not url.strip():
            continue
        prefix = f"[{sid}] " if sid is not None else ""
        if title:
            lines.append(f"- {prefix}{title} - {url}")
        else:
            lines.append(f"- {prefix}{url}")

    return "\n".join(lines).rstrip() + "\n"


def format_radar_text(radar: RadarReport) -> str:
    def _t(v: Any) -> str:
        it = _ensure_item(v)
        if it is None:
            return str(v).strip() if isinstance(v, str) else ""
        return str(it.get("text") or "").strip()

    def _score(v: Any) -> str:
        if isinstance(v, dict):
            s = v.get("score")
            try:
                return str(int(s))
            except Exception:
                return ""
        try:
            return str(int(v))
        except Exception:
            return ""

    lines: list[str] = []
    lines.append("## 赛道雷达（候选赛道）")
    lines.append("")
    lines.append("| # | 赛道 | 人群 | 变现（优先） | 难度(1-5) | 关键词(示例) | 风险(示例) |")
    lines.append("|---:|---|---|---|---:|---|---|")

    for i, niche in enumerate(radar.niches, start=1):
        name = str(niche.get("name") or "").strip()
        persona = _t(niche.get("target_persona"))

        monetization = niche.get("monetization") if isinstance(niche.get("monetization"), dict) else {}
        offers = _ensure_items((monetization or {}).get("offers"))
        offer_txt = " / ".join([it.get("text", "") for it in offers[:2] if it.get("text")])

        difficulty = _score(niche.get("difficulty"))
        kws = _ensure_items(niche.get("seo_keywords"))
        kw_txt = "、".join([it.get("text", "") for it in kws[:5] if it.get("text")])
        risks = _ensure_items(niche.get("compliance_risks"))
        risk_txt = "、".join([it.get("text", "") for it in risks[:2] if it.get("text")])

        lines.append(f"| {i} | {name} | {persona} | {offer_txt} | {difficulty} | {kw_txt} | {risk_txt} |")

    if radar.top3:
        lines.append("")
        lines.append("## Top3 推荐（先验证）")
        for x in radar.top3:
            name = str(x.get("name") or "").strip()
            why = x.get("why") if isinstance(x.get("why"), dict) else {}
            why_txt = _t(why)
            cites = _ensure_cites(why.get("cites"))
            suffix = f" [{','.join(str(i) for i in cites)}]" if cites else ""
            lines.append(f"- {name}：{why_txt}{suffix}")

    if radar.how_to_validate:
        lines.append("")
        lines.append("## 怎么验证（7 天内）")
        for it in radar.how_to_validate:
            lines.append(f"- {it.get('text','')}")

    lines.append("")
    lines.append("## 数据来源")
    for s in radar.sources:
        sid = s.get("id")
        url = s.get("url")
        title = s.get("title") or ""
        if not isinstance(url, str) or not url.strip():
            continue
        prefix = f"[{sid}] " if sid is not None else ""
        lines.append(f"- {prefix}{title} - {url}" if title else f"- {prefix}{url}")

    return "\n".join(lines).rstrip() + "\n"


def format_rewrite_text(rewrite: RewriteReport, *, sources: list[dict[str, Any]] | None = None) -> str:
    lines: list[str] = []
    lines.append("## 爆款仿写（Rewrite）")
    lines.append("")

    for d in rewrite.drafts:
        variant = d.get("variant") or ""
        tag = f"（{variant}）" if variant else ""
        cites = d.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"### 草稿{tag}: {d.get('title','')}{suffix}")
        lines.append(str(d.get("body") or "").rstrip())
        lines.append("")

    lines.append("### 标题库")
    for it in rewrite.title_bank:
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"- {it.get('text','')}{suffix}")
    lines.append("")

    lines.append("### 开头钩子")
    for it in rewrite.hooks:
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"- {it.get('text','')}{suffix}")
    lines.append("")

    lines.append("### 标签（Hashtags）")
    for it in rewrite.hashtags:
        lines.append(f"- {it.get('text','')}")
    lines.append("")

    lines.append("### 结尾 CTA")
    for it in rewrite.cta:
        lines.append(f"- {it.get('text','')}")
    lines.append("")

    lines.append("### 合规/避雷提示")
    for it in rewrite.compliance_notes:
        cites = it.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"- {it.get('text','')}{suffix}")
    lines.append("")

    if sources:
        lines.append("## 数据来源")
        for s in sources:
            sid = s.get("id")
            url = s.get("url")
            title = s.get("title") or ""
            if not isinstance(url, str) or not url.strip():
                continue
            prefix = f"[{sid}] " if sid is not None else ""
            lines.append(f"- {prefix}{title} - {url}" if title else f"- {prefix}{url}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def format_cover_text(cover: CoverReport, *, sources: list[dict[str, Any]] | None = None) -> str:
    lines: list[str] = []
    lines.append("## 封面导演（Cover Director）")
    lines.append("")

    for i, c in enumerate(cover.cover_concepts, start=1):
        name = c.get("name") or f"方案{i}"
        cites = c.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"### {name}{suffix}")
        lines.append(f"- 主标题：{c.get('headline','')}")
        lines.append(f"- 副标题：{c.get('subheadline','')}")
        mood = c.get("mood") or []
        if mood:
            lines.append(f"- 情绪：{', '.join(mood)}")
        if c.get("composition"):
            lines.append(f"- 构图：{c.get('composition')}")
        if c.get("lighting"):
            lines.append(f"- 光线：{c.get('lighting')}")
        props = c.get("props") or []
        if props:
            lines.append(f"- 道具：{', '.join(props)}")
        colors = c.get("color_palette") or []
        if colors:
            lines.append(f"- 配色：{', '.join(colors)}")
        if c.get("typography"):
            lines.append(f"- 字体/层级：{c.get('typography')}")
        if c.get("post_process"):
            lines.append(f"- 后期：{c.get('post_process')}")
        lines.append("")

    lines.append("### 拍摄清单（Shotlist）")
    for it in cover.shotlist:
        lines.append(f"- {it.get('text','')}")
    lines.append("")

    lines.append("### Canva/PS 操作步骤")
    for it in cover.canva_recipe:
        lines.append(f"- {it.get('text','')}")
    lines.append("")

    if cover.image_prompt is not None:
        cites = cover.image_prompt.get("cites") or []
        suffix = f" [{','.join(str(x) for x in cites)}]" if cites else ""
        lines.append(f"### 出图提示词（可选）{suffix}")
        lines.append(str(cover.image_prompt.get("text") or "").strip())
        lines.append("")

    if sources:
        lines.append("## 数据来源")
        for s in sources:
            sid = s.get("id")
            url = s.get("url")
            title = s.get("title") or ""
            if not isinstance(url, str) or not url.strip():
                continue
            prefix = f"[{sid}] " if sid is not None else ""
            lines.append(f"- {prefix}{title} - {url}" if title else f"- {prefix}{url}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def write_markdown(path: str | Path, *, query: str, report: ParsedReport) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    lines.append("# XHS Agent Report")
    lines.append("")
    lines.append(f"- Query: {query}")
    lines.append(f"- Generated: {_dt.datetime.now().isoformat(timespec='seconds')}")
    lines.append("")

    lines.append(format_report_text(report).rstrip())
    lines.append("")

    p.write_text("\n".join(lines), encoding="utf-8-sig")
    return p


def write_radar_markdown(path: str | Path, *, query: str, radar: RadarReport) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    lines.append("# XHS Niche Radar")
    lines.append("")
    lines.append(f"- Query: {query}")
    lines.append(f"- Generated: {_dt.datetime.now().isoformat(timespec='seconds')}")
    lines.append("")
    lines.append(format_radar_text(radar).rstrip())
    lines.append("")

    p.write_text("\n".join(lines), encoding="utf-8-sig")
    return p


def write_rewrite_markdown(
    path: str | Path, *, query: str, rewrite: RewriteReport, sources: list[dict[str, Any]] | None = None
) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    lines.append("# XHS Rewrite")
    lines.append("")
    lines.append(f"- Query: {query}")
    lines.append(f"- Generated: {_dt.datetime.now().isoformat(timespec='seconds')}")
    lines.append("")
    lines.append(format_rewrite_text(rewrite, sources=sources).rstrip())
    lines.append("")

    p.write_text("\n".join(lines), encoding="utf-8-sig")
    return p


def write_cover_markdown(
    path: str | Path, *, query: str, cover: CoverReport, sources: list[dict[str, Any]] | None = None
) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    lines.append("# XHS Cover Director")
    lines.append("")
    lines.append(f"- Query: {query}")
    lines.append(f"- Generated: {_dt.datetime.now().isoformat(timespec='seconds')}")
    lines.append("")
    lines.append(format_cover_text(cover, sources=sources).rstrip())
    lines.append("")

    p.write_text("\n".join(lines), encoding="utf-8-sig")
    return p


def write_csv_bundle(out_dir: str | Path, *, stem: str, report: ParsedReport) -> list[Path]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []

    def _write(name: str, rows: list[list[str]], header: list[str]) -> None:
        fp = d / name
        with fp.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
        written.append(fp)

    why = report.why_hot or {}
    st = report.structure or {}
    ideas = report.ideas or {}

    _write(
        f"{stem}_why_hot.csv",
        [["core_pain_points", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(why.get("core_pain_points"))]
        + [["emotional_hooks", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(why.get("emotional_hooks"))]
        + [["persona", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(why.get("persona"))],
        ["type", "text", "cites"],
    )

    _write(
        f"{stem}_structure.csv",
        [["title_templates", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(st.get("title_templates"))]
        + (
            [[
                "visual_style",
                (_ensure_item(st.get("visual_style")) or {}).get("text", ""),
                ",".join(str(x) for x in ((_ensure_item(st.get("visual_style")) or {}).get("cites") or [])),
            ]]
            if _ensure_item(st.get("visual_style")) is not None
            else []
        )
        + [["seo_keywords", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(st.get("seo_keywords"))],
        ["type", "text", "cites"],
    )

    _write(
        f"{stem}_ideas.csv",
        [["follow", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(ideas.get("follow"))]
        + [["reverse", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(ideas.get("reverse"))]
        + [["upgrade", it["text"], ",".join(str(x) for x in (it.get("cites") or []))] for it in _ensure_items(ideas.get("upgrade"))],
        ["angle", "text", "cites"],
    )

    _write(
        f"{stem}_sources.csv",
        [[str(s.get("id") or ""), str(s.get("url") or ""), str(s.get("title") or "")] for s in report.sources],
        ["id", "url", "title"],
    )

    return written


def write_rewrite_csv_bundle(
    out_dir: str | Path,
    *,
    rewrite: RewriteReport,
    sources: list[dict[str, Any]] | None = None,
) -> list[Path]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []

    def _write(name: str, rows: list[list[str]], header: list[str]) -> None:
        fp = d / name
        with fp.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
        written.append(fp)

    _write(
        "rewrite_drafts.csv",
        [
            [
                str(x.get("variant") or ""),
                str(x.get("title") or ""),
                str(x.get("body") or ""),
                ",".join(str(i) for i in (x.get("cites") or [])),
            ]
            for x in rewrite.drafts
        ],
        ["variant", "title", "body", "cites"],
    )

    _write(
        "rewrite_title_bank.csv",
        [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in rewrite.title_bank],
        ["text", "cites"],
    )
    _write(
        "rewrite_hooks.csv",
        [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in rewrite.hooks],
        ["text", "cites"],
    )
    _write("rewrite_hashtags.csv", [[it.get("text", "")] for it in rewrite.hashtags], ["text"])
    _write("rewrite_cta.csv", [[it.get("text", "")] for it in rewrite.cta], ["text"])
    _write(
        "rewrite_compliance_notes.csv",
        [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in rewrite.compliance_notes],
        ["text", "cites"],
    )

    if sources is not None:
        _write(
            "rewrite_sources.csv",
            [[str(s.get("id") or ""), str(s.get("url") or ""), str(s.get("title") or "")] for s in sources],
            ["id", "url", "title"],
        )

    return written


def write_cover_csv_bundle(
    out_dir: str | Path,
    *,
    cover: CoverReport,
    sources: list[dict[str, Any]] | None = None,
) -> list[Path]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []

    def _write(name: str, rows: list[list[str]], header: list[str]) -> None:
        fp = d / name
        with fp.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
        written.append(fp)

    _write(
        "cover_concepts.csv",
        [
            [
                str(c.get("name") or ""),
                str(c.get("headline") or ""),
                str(c.get("subheadline") or ""),
                ";".join(str(x) for x in (c.get("mood") or [])),
                str(c.get("composition") or ""),
                str(c.get("lighting") or ""),
                ";".join(str(x) for x in (c.get("props") or [])),
                ";".join(str(x) for x in (c.get("color_palette") or [])),
                str(c.get("typography") or ""),
                str(c.get("post_process") or ""),
                ",".join(str(i) for i in (c.get("cites") or [])),
            ]
            for c in cover.cover_concepts
        ],
        [
            "name",
            "headline",
            "subheadline",
            "mood",
            "composition",
            "lighting",
            "props",
            "color_palette",
            "typography",
            "post_process",
            "cites",
        ],
    )

    _write(
        "cover_shotlist.csv",
        [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in cover.shotlist],
        ["text", "cites"],
    )
    _write(
        "cover_canva_recipe.csv",
        [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in cover.canva_recipe],
        ["text", "cites"],
    )
    if cover.image_prompt is not None:
        _write(
            "cover_image_prompt.csv",
            [[cover.image_prompt.get("text", ""), ",".join(str(i) for i in (cover.image_prompt.get("cites") or []))]],
            ["text", "cites"],
        )

    if sources is not None:
        _write(
            "cover_sources.csv",
            [[str(s.get("id") or ""), str(s.get("url") or ""), str(s.get("title") or "")] for s in sources],
            ["id", "url", "title"],
        )

    return written


def write_radar_csv_bundle(out_dir: str | Path, *, radar: RadarReport) -> list[Path]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []

    def _write(name: str, rows: list[list[str]], header: list[str]) -> None:
        fp = d / name
        with fp.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)
        written.append(fp)

    def _t(v: Any) -> str:
        it = _ensure_item(v)
        if it is None:
            return str(v).strip() if isinstance(v, str) else ""
        return str(it.get("text") or "").strip()

    def _items(v: Any, *, limit: int = 10) -> str:
        return " | ".join([it.get("text", "") for it in _ensure_items(v)[:limit] if it.get("text")])

    def _difficulty(v: Any) -> str:
        if isinstance(v, dict) and v.get("score") is not None:
            try:
                return str(int(v.get("score")))
            except Exception:
                return str(v.get("score"))
        return str(v or "")

    rows: list[list[str]] = []
    for niche in radar.niches:
        monetization = niche.get("monetization") if isinstance(niche.get("monetization"), dict) else {}
        rows.append(
            [
                str(niche.get("name") or "").strip(),
                _t(niche.get("target_persona")),
                _items(niche.get("core_pain_points"), limit=6),
                _items(niche.get("content_pillars"), limit=6),
                _items(niche.get("seo_keywords"), limit=10),
                _items(niche.get("title_templates"), limit=6),
                _items(niche.get("content_structures"), limit=6),
                _items((monetization or {}).get("offers"), limit=6),
                _items((monetization or {}).get("delivery"), limit=6),
                _t((monetization or {}).get("first_offer_hint")),
                _difficulty(niche.get("difficulty")),
                _items(niche.get("compliance_risks"), limit=6),
            ]
        )

    _write(
        "radar_niches.csv",
        rows,
        [
            "name",
            "target_persona",
            "core_pain_points",
            "content_pillars",
            "seo_keywords",
            "title_templates",
            "content_structures",
            "monetization_offers",
            "monetization_delivery",
            "first_offer_hint",
            "difficulty",
            "compliance_risks",
        ],
    )

    if radar.top3:
        top_rows: list[list[str]] = []
        for x in radar.top3:
            why = x.get("why") if isinstance(x.get("why"), dict) else {}
            top_rows.append([str(x.get("name") or "").strip(), _t(why), ",".join(str(i) for i in _ensure_cites(why.get("cites")))])
        _write("radar_top3.csv", top_rows, ["name", "why", "cites"])

    if radar.how_to_validate:
        _write(
            "radar_how_to_validate.csv",
            [[it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))] for it in radar.how_to_validate],
            ["text", "cites"],
        )

    _write(
        "radar_sources.csv",
        [[str(s.get("id") or ""), str(s.get("url") or ""), str(s.get("title") or "")] for s in radar.sources],
        ["id", "url", "title"],
    )

    return written


def write_json(path: str | Path, *, data: Any) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


def write_xlsx(path: str | Path, *, report: ParsedReport) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as e:  # pragma: no cover
        raise ImportError("openpyxl is required to write .xlsx. Install: pip install openpyxl") from e

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    why = report.why_hot or {}
    st = report.structure or {}
    ideas = report.ideas or {}

    ws = wb.active
    ws.title = "why_hot"
    ws.append(["type", "text", "cites"])
    for it in _ensure_items(why.get("core_pain_points")):
        ws.append(["core_pain_points", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])
    for it in _ensure_items(why.get("emotional_hooks")):
        ws.append(["emotional_hooks", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])
    for it in _ensure_items(why.get("persona")):
        ws.append(["persona", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])

    ws = wb.create_sheet("structure")
    ws.append(["type", "text", "cites"])
    for it in _ensure_items(st.get("title_templates")):
        ws.append(["title_templates", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])
    vs_it = _ensure_item(st.get("visual_style"))
    if vs_it is not None:
        ws.append(["visual_style", vs_it.get("text", ""), ",".join(str(x) for x in (vs_it.get("cites") or []))])
    for it in _ensure_items(st.get("seo_keywords")):
        ws.append(["seo_keywords", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])

    ws = wb.create_sheet("ideas")
    ws.append(["angle", "text", "cites"])
    for it in _ensure_items(ideas.get("follow")):
        ws.append(["follow", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])
    for it in _ensure_items(ideas.get("reverse")):
        ws.append(["reverse", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])
    for it in _ensure_items(ideas.get("upgrade")):
        ws.append(["upgrade", it["text"], ",".join(str(x) for x in (it.get("cites") or []))])

    ws = wb.create_sheet("sources")
    ws.append(["id", "url", "title"])
    for s in report.sources:
        ws.append([s.get("id"), s.get("url"), s.get("title")])

    wb.save(p)
    return p


def write_radar_xlsx(path: str | Path, *, radar: RadarReport) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as e:  # pragma: no cover
        raise ImportError("openpyxl is required to write .xlsx. Install: pip install openpyxl") from e

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    def _t(v: Any) -> str:
        it = _ensure_item(v)
        if it is None:
            return str(v).strip() if isinstance(v, str) else ""
        return str(it.get("text") or "").strip()

    def _items(v: Any, *, limit: int = 10) -> str:
        return " | ".join([it.get("text", "") for it in _ensure_items(v)[:limit] if it.get("text")])

    def _difficulty(v: Any) -> str:
        if isinstance(v, dict) and v.get("score") is not None:
            try:
                return str(int(v.get("score")))
            except Exception:
                return str(v.get("score"))
        return str(v or "")

    wb = Workbook()
    ws = wb.active
    ws.title = "niches"
    ws.append(
        [
            "name",
            "target_persona",
            "core_pain_points",
            "content_pillars",
            "seo_keywords",
            "title_templates",
            "content_structures",
            "monetization_offers",
            "monetization_delivery",
            "first_offer_hint",
            "difficulty",
            "compliance_risks",
        ]
    )

    for niche in radar.niches:
        monetization = niche.get("monetization") if isinstance(niche.get("monetization"), dict) else {}
        ws.append(
            [
                str(niche.get("name") or "").strip(),
                _t(niche.get("target_persona")),
                _items(niche.get("core_pain_points"), limit=6),
                _items(niche.get("content_pillars"), limit=6),
                _items(niche.get("seo_keywords"), limit=10),
                _items(niche.get("title_templates"), limit=6),
                _items(niche.get("content_structures"), limit=6),
                _items((monetization or {}).get("offers"), limit=6),
                _items((monetization or {}).get("delivery"), limit=6),
                _t((monetization or {}).get("first_offer_hint")),
                _difficulty(niche.get("difficulty")),
                _items(niche.get("compliance_risks"), limit=6),
            ]
        )

    ws = wb.create_sheet("top3")
    ws.append(["name", "why", "cites"])
    for x in radar.top3:
        why = x.get("why") if isinstance(x.get("why"), dict) else {}
        ws.append([str(x.get("name") or "").strip(), _t(why), ",".join(str(i) for i in _ensure_cites(why.get("cites")))])

    ws = wb.create_sheet("how_to_validate")
    ws.append(["text", "cites"])
    for it in radar.how_to_validate:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    ws = wb.create_sheet("sources")
    ws.append(["id", "url", "title"])
    for s in radar.sources:
        ws.append([s.get("id"), s.get("url"), s.get("title")])

    wb.save(p)
    return p


def write_rewrite_xlsx(path: str | Path, *, rewrite: RewriteReport, sources: list[dict[str, Any]] | None = None) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as e:  # pragma: no cover
        raise ImportError("openpyxl is required to write .xlsx. Install: pip install openpyxl") from e

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws = wb.active
    ws.title = "drafts"
    ws.append(["variant", "title", "body", "cites"])
    for d in rewrite.drafts:
        ws.append(
            [
                d.get("variant") or "",
                d.get("title") or "",
                d.get("body") or "",
                ",".join(str(i) for i in (d.get("cites") or [])),
            ]
        )

    ws = wb.create_sheet("title_bank")
    ws.append(["text", "cites"])
    for it in rewrite.title_bank:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    ws = wb.create_sheet("hooks")
    ws.append(["text", "cites"])
    for it in rewrite.hooks:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    ws = wb.create_sheet("hashtags")
    ws.append(["text"])
    for it in rewrite.hashtags:
        ws.append([it.get("text", "")])

    ws = wb.create_sheet("cta")
    ws.append(["text"])
    for it in rewrite.cta:
        ws.append([it.get("text", "")])

    ws = wb.create_sheet("compliance_notes")
    ws.append(["text", "cites"])
    for it in rewrite.compliance_notes:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    if sources is not None:
        ws = wb.create_sheet("sources")
        ws.append(["id", "url", "title"])
        for s in sources:
            ws.append([s.get("id"), s.get("url"), s.get("title")])

    wb.save(p)
    return p


def write_cover_xlsx(path: str | Path, *, cover: CoverReport, sources: list[dict[str, Any]] | None = None) -> Path:
    try:
        from openpyxl import Workbook  # type: ignore
    except Exception as e:  # pragma: no cover
        raise ImportError("openpyxl is required to write .xlsx. Install: pip install openpyxl") from e

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws = wb.active
    ws.title = "cover_concepts"
    ws.append(
        [
            "name",
            "headline",
            "subheadline",
            "mood",
            "composition",
            "lighting",
            "props",
            "color_palette",
            "typography",
            "post_process",
            "cites",
        ]
    )
    for c in cover.cover_concepts:
        ws.append(
            [
                c.get("name") or "",
                c.get("headline") or "",
                c.get("subheadline") or "",
                ";".join(str(x) for x in (c.get("mood") or [])),
                c.get("composition") or "",
                c.get("lighting") or "",
                ";".join(str(x) for x in (c.get("props") or [])),
                ";".join(str(x) for x in (c.get("color_palette") or [])),
                c.get("typography") or "",
                c.get("post_process") or "",
                ",".join(str(i) for i in (c.get("cites") or [])),
            ]
        )

    ws = wb.create_sheet("shotlist")
    ws.append(["text", "cites"])
    for it in cover.shotlist:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    ws = wb.create_sheet("canva_recipe")
    ws.append(["text", "cites"])
    for it in cover.canva_recipe:
        ws.append([it.get("text", ""), ",".join(str(i) for i in (it.get("cites") or []))])

    if cover.image_prompt is not None:
        ws = wb.create_sheet("image_prompt")
        ws.append(["text", "cites"])
        ws.append(
            [
                cover.image_prompt.get("text", ""),
                ",".join(str(i) for i in (cover.image_prompt.get("cites") or [])),
            ]
        )

    if sources is not None:
        ws = wb.create_sheet("sources")
        ws.append(["id", "url", "title"])
        for s in sources:
            ws.append([s.get("id"), s.get("url"), s.get("title")])

    wb.save(p)
    return p
